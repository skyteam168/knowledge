"""Sample document generators shared by parser / pipeline tests and acceptance runs."""

import io

LOREM_ZH = (
    "企业知识库平台用于统一管理制度文件、技术文档与业务资料。"
    "文档上传后会经过解析、切分、向量化与索引等步骤，最终支持混合检索与智能问答。"
    "为保证安全，所有检索结果都会按照用户的部门、用户组与角色进行权限过滤。"
)
LOREM_EN = (
    "The ingestion pipeline parses each document into a section tree and splits it into "
    "parent and child chunks. Child chunks are the retrieval units, while parent chunks "
    "provide the surrounding context for answer generation. "
)

TABLE_HEADER = ["编号", "物料名称", "规格", "单价(元)", "库存"]


def table_rows(count: int) -> list[list[str]]:
    return [
        [f"M-{i:04d}", f"飞织鞋面材料{i}", f"{100 + i}g/m²", f"{12 + i % 7}.50", str(i * 3)]
        for i in range(1, count + 1)
    ]


def make_docx(table_rows_count: int = 60) -> bytes:
    import docx

    document = docx.Document()
    document.core_properties.title = "员工手册"
    document.add_heading("第一章 总则", level=1)
    document.add_paragraph(LOREM_ZH * 2)
    document.add_heading("1.1 适用范围", level=2)
    document.add_paragraph(LOREM_ZH * 6)
    document.add_heading("1.2 术语", level=2)
    document.add_paragraph(LOREM_ZH)
    document.add_heading("第二章 物料清单", level=1)
    document.add_paragraph("下表列出了常用物料：")
    table = document.add_table(rows=1, cols=len(TABLE_HEADER))
    for cell, text in zip(table.rows[0].cells, TABLE_HEADER, strict=True):
        cell.text = text
    for row in table_rows(table_rows_count):
        cells = table.add_row().cells
        for cell, text in zip(cells, row, strict=True):
            cell.text = text
    document.add_heading("第三章 附则", level=1)
    document.add_paragraph("本手册自发布之日起施行。")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def make_xlsx(rows: int = 120) -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "库存"
    sheet.append(TABLE_HEADER)
    for row in table_rows(rows):
        sheet.append(row)
    summary = workbook.create_sheet("汇总")
    summary.append(["类别", "数量"])
    summary.append(["鞋面", 42])
    summary.append(["鞋底", 17])
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


def make_pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "季度经营回顾"
    slide.placeholders[1].text = "营收同比增长 18%\n新增客户 32 家"
    slide.notes_slide.notes_text_frame.text = "强调华东区贡献最大。"

    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "区域销售数据"
    rows = [["区域", "销售额", "同比"], ["华东", "1200", "+20%"], ["华南", "860", "+12%"]]
    shape = slide.shapes.add_table(len(rows), 3, Inches(1), Inches(1.5), Inches(6), Inches(2))
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            shape.table.cell(r, c).text = text

    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "下季度计划"
    slide.placeholders[1].text = "推进飞织新品上市\n完善知识库建设"
    buf = io.BytesIO()
    presentation.save(buf)
    return buf.getvalue()


def make_pdf(font_path: str | None = None) -> bytes:
    """Headings in larger fonts, body text and a bordered table. CJK needs a TTF font."""
    from fpdf import FPDF

    pdf = FPDF()
    family = "Helvetica"
    if font_path:
        pdf.add_font("cjk", "", font_path)
        pdf.add_font("cjk", "B", font_path)
        family = "cjk"
    pdf.set_title("Ingestion Guide")
    body = (LOREM_ZH if font_path else LOREM_EN) * 3

    pdf.add_page()
    pdf.set_font(family, size=20)
    pdf.cell(0, 12, "1 Overview" if not font_path else "1 概述", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font(family, size=11)
    pdf.multi_cell(0, 6, body)
    pdf.set_font(family, size=16)
    pdf.cell(0, 10, "1.1 Scope" if not font_path else "1.1 范围", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font(family, size=11)
    pdf.multi_cell(0, 6, body)

    pdf.add_page()
    pdf.set_font(family, size=20)
    pdf.cell(0, 12, "2 Prices" if not font_path else "2 价格表", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font(family, size=10)
    header = ["Code", "Name", "Spec", "Price", "Stock"] if not font_path else TABLE_HEADER
    rows = [
        [f"M-{i:04d}", f"Material {i}", f"{100 + i}g", f"{12 + i % 7}.50", str(i * 3)]
        for i in range(1, 16)
    ]
    with pdf.table() as table:
        table.row(header)
        for row in rows:
            table.row(row)
    return bytes(pdf.output())


def make_markdown() -> bytes:
    text = f"""# 部署指南

## 环境要求

{LOREM_ZH}

## 配置项

| 名称 | 默认值 | 说明 |
| --- | --- | --- |
| PDF_PARSER | auto | PDF 解析后端 |
| OCR_ENABLED | true | 扫描件 OCR |

## 常见问题

{LOREM_ZH}
"""
    return text.encode("utf-8")
