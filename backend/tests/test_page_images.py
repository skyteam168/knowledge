"""Unit tests for PDF/PPT page raster extraction (citation previews)."""

from io import BytesIO

from PIL import Image

from app.services.ingest.page_images import asset_key, extract_page_images
from tests.samples import make_pdf, make_pptx


def _tiny_png() -> bytes:
    buf = BytesIO()
    Image.new("RGB", (40, 30), color=(20, 120, 200)).save(buf, format="PNG")
    return buf.getvalue()


def make_pptx_with_picture() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "带图幻灯片"
    slide.shapes.add_picture(BytesIO(_tiny_png()), Inches(1), Inches(1.5), width=Inches(3))
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "纯文字页"
    slide.placeholders[1].text = "没有嵌入图片"
    buf = BytesIO()
    presentation.save(buf)
    return buf.getvalue()


def test_asset_key_layout() -> None:
    assert asset_key("org/kb/doc/v1/file.pdf", "page", 3) == "org/kb/doc/v1/pages/3.png"
    assert asset_key("org/kb/doc/v1/file.pptx", "slide", 2) == "org/kb/doc/v1/slides/2.png"


def test_extract_pdf_page_images() -> None:
    images = extract_page_images(make_pdf(), "guide.pdf", dpi=72, max_edge=800)
    assert images
    assert all(i.kind == "page" for i in images)
    assert images[0].page_index == 1
    assert images[0].mime_type == "image/png"
    assert images[0].data[:8] == b"\x89PNG\r\n\x1a\n"
    assert images[0].width and images[0].width <= 800


def test_extract_pptx_falls_back_to_embedded_pictures() -> None:
    images = extract_page_images(make_pptx_with_picture(), "deck.pptx", dpi=72, max_edge=400)
    # Without LibreOffice, only slides that contain pictures produce assets.
    assert len(images) == 1
    assert images[0].kind == "slide"
    assert images[0].page_index == 1
    assert images[0].mime_type == "image/png"


def test_extract_pptx_without_pictures_returns_empty_without_soffice() -> None:
    images = extract_page_images(make_pptx(), "review.pptx", dpi=72)
    # Text-only deck + no soffice → nothing to show (retrieval still works on text).
    assert images == []


def test_extract_raster_image_as_single_page_asset() -> None:
    images = extract_page_images(_tiny_png(), "photo.png", max_edge=200)
    assert len(images) == 1
    assert images[0].kind == "page"
    assert images[0].page_index == 1
    assert images[0].mime_type == "image/png"
    assert images[0].width and images[0].width <= 200
