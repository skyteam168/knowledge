"""Excel import: one row per (person, department); a person in several departments has
several rows. Department paths like "集团/越南Flyknit/织造车间" create the whole chain;
each path prefix is the department's external id.
"""

import io
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel

from app.core.exceptions import ValidationAppError
from app.services.directory.connectors.base import DirectoryConnector
from app.services.directory.snapshot import DeptRecord, Snapshot, UserRecord

COLUMNS = ["部门路径", "账号", "工号", "姓名", "邮箱", "手机", "职务", "是否主部门", "在职"]
REQUIRED = ["部门路径", "账号", "姓名"]
TRUE_VALUES = {"是", "y", "yes", "true", "1", "在职", "√"}
FALSE_VALUES = {"否", "n", "no", "false", "0", "离职", "×"}
MAX_ROWS = 50_000
EXAMPLES = [
    [
        "集团/越南Flyknit/织造车间",
        "zhangsan",
        "VN0001",
        "张三",
        "zhangsan@example.com",
        "",
        "班组长",
        "是",
        "是",
    ],
    [
        "集团/越南Flyknit/品质部",
        "zhangsan",
        "VN0001",
        "张三",
        "zhangsan@example.com",
        "",
        "班组长",
        "否",
        "是",
    ],
    ["集团/总部/法务部", "lisi", "HQ0002", "李四", "", "13800000000", "法务专员", "是", "是"],
]


class ExcelConfig(BaseModel):
    """Excel sources keep no connection settings; data comes with each upload."""


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _flag(value: str, default: bool) -> bool:
    text = value.strip().lower()
    if text in TRUE_VALUES:
        return True
    if text in FALSE_VALUES:
        return False
    return default


def build_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "人员"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        if cell.value in REQUIRED:
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
    for row in EXAMPLES:
        ws.append(row)
    for col, width in zip("ABCDEFGHI", [36, 16, 12, 12, 26, 16, 14, 12, 8], strict=True):
        ws.column_dimensions[col].width = width
    help_ws = wb.create_sheet("说明")
    for line in [
        ["列", "说明"],
        ["部门路径", "必填。用 / 分隔的完整路径，如：集团/越南Flyknit/织造车间；上级部门自动创建"],
        ["账号", "必填。登录账号（与 Keycloak / 企业微信账号一致时可直接扫码或密码登录）"],
        ["工号", "选填。用于与已有账号匹配（优先级：工号 > 手机 > 邮箱 > 账号）"],
        ["姓名", "必填"],
        ["邮箱 / 手机", "选填"],
        ["职务", "选填。会生成用户组「职务/<职务>」，可直接用于知识库授权"],
        ["是否主部门", "一人多部门时写多行，标记「是」的为主部门；不填则第一行为主部门"],
        ["在职", "是 / 否，默认是。表格里没有的人会被停用（预览时可以先检查）"],
    ]:
        help_ws.append(line)
    help_ws.column_dimensions["A"].width = 14
    help_ws.column_dimensions["B"].width = 90
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_workbook(content: bytes) -> Snapshot:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationAppError("无法读取 Excel 文件，请使用模板（.xlsx）") from exc
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [_cell(v) for v in next(rows, ())]
    missing = [c for c in REQUIRED if c not in header]
    if missing:
        raise ValidationAppError("Excel 缺少必填列", detail={"missing": missing})
    pos = {name: header.index(name) for name in COLUMNS if name in header}

    def get(row: tuple[Any, ...], name: str) -> str:
        idx = pos.get(name)
        return _cell(row[idx]) if idx is not None and idx < len(row) else ""

    depts: dict[str, DeptRecord] = {}
    users: dict[str, UserRecord] = {}
    errors: list[str] = []
    for line_no, row in enumerate(rows, start=2):
        if line_no > MAX_ROWS + 1:
            raise ValidationAppError(f"最多 {MAX_ROWS} 行")
        if not any(_cell(v) for v in row):
            continue
        path = get(row, "部门路径")
        username = get(row, "账号")
        name = get(row, "姓名")
        if not (path and username and name):
            errors.append(f"第 {line_no} 行缺少部门路径 / 账号 / 姓名")
            continue
        parts = [p.strip() for p in path.replace("\\", "/").split("/") if p.strip()]
        parent: str | None = None
        for i in range(len(parts)):
            ext = "/".join(parts[: i + 1])
            if ext not in depts:
                depts[ext] = DeptRecord(
                    external_id=ext, name=parts[i], parent_external_id=parent, order=len(depts)
                )
            parent = ext
        dept_ext = "/".join(parts)
        user = users.get(username)
        if user is None:
            user = users[username] = UserRecord(
                external_id=username,
                username=username,
                name=name,
                employee_no=get(row, "工号") or None,
                email=get(row, "邮箱") or None,
                mobile=get(row, "手机") or None,
                title=get(row, "职务") or None,
                active=_flag(get(row, "在职"), True),
            )
        if dept_ext not in user.dept_external_ids:
            user.dept_external_ids.append(dept_ext)
        if _flag(get(row, "是否主部门"), False):
            user.primary_dept = dept_ext
    wb.close()
    if errors:
        raise ValidationAppError("Excel 数据有误", detail={"errors": errors[:50]})
    return Snapshot(departments=list(depts.values()), users=list(users.values()))


class ExcelConnector(DirectoryConnector):
    config_model = ExcelConfig

    def __init__(self, config: dict[str, Any], settings: Any, content: bytes | None = None) -> None:
        super().__init__(config, settings)
        self.content = content

    async def fetch(self) -> Snapshot:
        if self.content is None:
            raise ValidationAppError("Excel 数据源需要上传文件", code="DIRECTORY_UPLOAD_REQUIRED")
        return parse_workbook(self.content)
