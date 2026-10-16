import csv
import io
import re
import zipfile
from datetime import date, datetime, time
from typing import Any

from app.services.ingest.parsers.base import ParseError, decode_text, stem
from app.services.ingest.sections import (
    Location,
    ParsedDocument,
    SectionBuilder,
    TableData,
    table_from_rows,
)

_HEADING_STYLE = re.compile(r"^(?:heading|标题)\s*(\d)$", re.IGNORECASE)


def _open_error(exc: Exception, kind: str) -> ParseError:
    return ParseError(f"Cannot open {kind} file: {type(exc).__name__}: {exc}")


_OLE2 = b"\xd0\xcf\x11\xe0"


class WordParser:
    """.docx directly; legacy .doc is converted to .docx via LibreOffice when available."""

    name = "docx"

    def __init__(self, soffice_path: str | None = None) -> None:
        self.soffice_path = soffice_path

    def parse(self, data: bytes, filename: str) -> ParsedDocument:
        import docx
        from docx.document import Document as DocxDocument
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        payload, parser_name, meta = self._normalize(data, filename)
        try:
            document: DocxDocument = docx.Document(io.BytesIO(payload))
        except (zipfile.BadZipFile, KeyError, ValueError) as exc:
            raise _open_error(exc, "Word") from exc

        title = (document.core_properties.title or "").strip() or stem(filename)
        builder = SectionBuilder(title)
        for element in document.element.body.iterchildren():
            tag = element.tag.rsplit("}", 1)[-1]
            if tag == "p":
                paragraph = Paragraph(element, document)
                text = paragraph.text.strip()
                if not text:
                    continue
                level = self._heading_level(paragraph)
                if level is not None:
                    builder.heading(text, level)
                else:
                    builder.text(text)
            elif tag == "tbl":
                table = Table(element, document)
                rows = [self._row_cells(row) for row in table.rows]
                parsed = table_from_rows(rows)
                if parsed:
                    builder.table(parsed)
        parsed_doc = builder.build(parser_name)
        if meta:
            parsed_doc.metadata.update(meta)
        return parsed_doc

    def _normalize(
        self, data: bytes, filename: str
    ) -> tuple[bytes, str, dict[str, Any]]:
        lower = filename.lower()
        legacy = lower.endswith(".doc") and not lower.endswith(".docx")
        ole2 = data.startswith(_OLE2)
        if not legacy and not ole2:
            return data, self.name, {}
        from app.services.ingest.office_convert import convert_with_soffice
        from app.services.ingest.parsers.base import ParseError

        converted = convert_with_soffice(
            data,
            source_suffix=".doc",
            to_format="docx",
            soffice_path=self.soffice_path,
        )
        if not converted:
            raise ParseError(
                "Legacy .doc requires LibreOffice (soffice) to convert to .docx. "
                "Install LibreOffice or re-save the file as .docx."
            )
        return converted, "doc", {"converted_from": "doc"}

    @staticmethod
    def _row_cells(row: Any) -> list[str]:
        cells: list[str] = []
        previous = None
        for cell in row.cells:
            # Horizontally merged cells are reported once per grid column.
            if previous is not None and cell._tc is previous:
                continue
            previous = cell._tc
            cells.append(cell.text.strip())
        return cells

    @staticmethod
    def _heading_level(paragraph: Any) -> int | None:
        style = paragraph.style
        while style is not None:
            name = (style.name or "").strip()
            if name.lower() == "title":
                return 1
            m = _HEADING_STYLE.match(name)
            if m:
                return int(m.group(1))
            style = style.base_style
        ppr = paragraph._p.pPr
        if ppr is not None:
            outline = ppr.find(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}outlineLvl"
            )
            if outline is not None:
                value = outline.get(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val"
                )
                if value is not None and value.isdigit() and int(value) < 9:
                    return int(value) + 1
        return None


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds").removesuffix(" 00:00:00")
    if isinstance(value, date | time):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _split_on_blank_rows(rows: list[list[str]]) -> list[list[list[str]]]:
    groups: list[list[list[str]]] = [[]]
    for row in rows:
        if any(row):
            groups[-1].append(row)
        elif groups[-1]:
            groups.append([])
    return [g for g in groups if g]


def _add_sheet_tables(builder: SectionBuilder, rows: list[list[str]], location: Location) -> None:
    for group in _split_on_blank_rows(rows):
        if len(group) == 1:
            builder.text("  ".join(c for c in group[0] if c), location)
            continue
        table = table_from_rows(group)
        if table:
            builder.table(table, location)


class ExcelParser:
    """.xlsx: one section per sheet; blank-row separated regions become separate tables."""

    name = "xlsx"

    def parse(self, data: bytes, filename: str) -> ParsedDocument:
        import openpyxl

        try:
            workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
            raise _open_error(exc, "Excel") from exc

        builder = SectionBuilder(stem(filename))
        sheets: list[str] = []
        try:
            for sheet in workbook.worksheets:
                location = Location(sheet=sheet.title)
                rows = [[_cell_text(v) for v in row] for row in sheet.iter_rows(values_only=True)]
                if not any(any(r) for r in rows):
                    continue
                sheets.append(sheet.title)
                builder.heading(sheet.title, 1, location)
                _add_sheet_tables(builder, rows, location)
        finally:
            workbook.close()
        return builder.build(self.name, {"sheets": sheets})


class CsvParser:
    name = "csv"

    def parse(self, data: bytes, filename: str) -> ParsedDocument:
        text = decode_text(data)
        try:
            dialect: Any = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
        builder = SectionBuilder(stem(filename))
        location = Location(sheet=stem(filename))
        _add_sheet_tables(builder, rows, location)
        return builder.build(self.name)


class PowerPointParser:
    """.pptx: one section per slide with body text, tables and speaker notes."""

    name = "pptx"

    def parse(self, data: bytes, filename: str) -> ParsedDocument:
        from pptx import Presentation

        try:
            presentation = Presentation(io.BytesIO(data))
        except (zipfile.BadZipFile, KeyError, ValueError) as exc:
            raise _open_error(exc, "PowerPoint") from exc

        title = (presentation.core_properties.title or "").strip() or stem(filename)
        builder = SectionBuilder(title)
        for number, slide in enumerate(presentation.slides, start=1):
            location = Location(slide=number)
            title_shape = slide.shapes.title
            slide_title = title_shape.text_frame.text.strip() if title_shape is not None else ""
            builder.heading(slide_title or f"第 {number} 页", 1, location)
            for shape in slide.shapes:
                if title_shape is not None and shape.shape_id == title_shape.shape_id:
                    continue
                self._shape(shape, builder, location)
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame
                if notes is not None and notes.text.strip():
                    builder.text(f"备注：{notes.text.strip()}", location)
        return builder.build(self.name, {"slides": len(presentation.slides)})

    def _shape(self, shape: Any, builder: SectionBuilder, location: Location) -> None:
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for inner in shape.shapes:
                self._shape(inner, builder, location)
            return
        if getattr(shape, "has_table", False) and shape.has_table:
            rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
            table: TableData | None = table_from_rows(rows)
            if table:
                builder.table(table, location)
            return
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            lines = [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
            if lines:
                builder.text("\n".join(lines), location)
