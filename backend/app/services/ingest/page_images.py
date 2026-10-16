"""Rasterize PDF pages / PPT slides into PNG assets for citation previews.

Images are stored beside the original object and linked by page/slide index.
They are never embedded into chunk text, so retrieval scoring is unchanged.
"""

from __future__ import annotations

import hashlib
import io
import logging
from dataclasses import dataclass
from typing import Any

logger = logging.getLogger(__name__)

DEFAULT_DPI = 144
DEFAULT_MAX_EDGE = 1600
DEFAULT_MAX_PAGES = 200


@dataclass(frozen=True)
class PageImage:
    kind: str  # PageAssetKind value
    page_index: int  # 1-based
    data: bytes
    mime_type: str
    width: int | None
    height: int | None
    content_hash: str


def asset_key(minio_key: str, kind: str, page_index: int) -> str:
    from posixpath import dirname

    folder = "pages" if kind == "page" else "slides"
    return f"{dirname(minio_key)}/{folder}/{page_index}.png"


def extract_page_images(
    data: bytes,
    filename: str,
    *,
    dpi: int = DEFAULT_DPI,
    max_edge: int = DEFAULT_MAX_EDGE,
    max_pages: int = DEFAULT_MAX_PAGES,
    soffice_path: str | None = None,
) -> list[PageImage]:
    """Best-effort page/slide rasters. Returns [] when unsupported or empty."""
    name = filename.lower()
    try:
        if name.endswith(".pdf"):
            return _from_pdf(data, dpi=dpi, max_edge=max_edge, max_pages=max_pages)
        if name.endswith(".pptx"):
            return _from_pptx(
                data,
                dpi=dpi,
                max_edge=max_edge,
                max_pages=max_pages,
                soffice_path=soffice_path,
            )
        if name.endswith((".png", ".jpg", ".jpeg")):
            return _from_raster(data, max_edge=max_edge)
    except Exception:
        logger.exception("page_image_extract_failed", extra={"filename": filename})
    return []


def _from_raster(data: bytes, *, max_edge: int) -> list[PageImage]:
    from PIL import Image

    image = Image.open(io.BytesIO(data))
    image.load()
    return [_page_image("page", 1, image, max_edge)]


def _encode_png(image: Any, max_edge: int) -> tuple[bytes, int, int]:
    from PIL import Image

    if not isinstance(image, Image.Image):
        image = Image.fromarray(image)
    image = image.convert("RGB")
    w, h = image.size
    longest = max(w, h)
    if longest > max_edge:
        scale = max_edge / longest
        image = image.resize(
            (max(1, int(w * scale)), max(1, int(h * scale))),
            Image.Resampling.LANCZOS,
        )
        w, h = image.size
    buf = io.BytesIO()
    image.save(buf, format="PNG", optimize=True)
    raw = buf.getvalue()
    return raw, w, h


def _page_image(kind: str, page_index: int, image: Any, max_edge: int) -> PageImage:
    data, width, height = _encode_png(image, max_edge)
    return PageImage(
        kind=kind,
        page_index=page_index,
        data=data,
        mime_type="image/png",
        width=width,
        height=height,
        content_hash=hashlib.sha256(data).hexdigest(),
    )


def _from_pdf(
    data: bytes, *, dpi: int, max_edge: int, max_pages: int, kind: str = "page"
) -> list[PageImage]:
    import pdfplumber

    out: list[PageImage] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for i, page in enumerate(pdf.pages[:max_pages], start=1):
            image = page.to_image(resolution=dpi).original
            out.append(_page_image(kind, i, image, max_edge))
    return out


def _from_pptx(
    data: bytes,
    *,
    dpi: int,
    max_edge: int,
    max_pages: int,
    soffice_path: str | None,
) -> list[PageImage]:
    pdf_bytes = _pptx_to_pdf(data, soffice_path=soffice_path)
    if pdf_bytes:
        return _from_pdf(pdf_bytes, dpi=dpi, max_edge=max_edge, max_pages=max_pages, kind="slide")
    return _pptx_embedded_pictures(data, max_edge=max_edge, max_pages=max_pages)


def _pptx_to_pdf(data: bytes, *, soffice_path: str | None) -> bytes | None:
    from app.services.ingest.office_convert import convert_with_soffice

    return convert_with_soffice(
        data, source_suffix=".pptx", to_format="pdf", soffice_path=soffice_path
    )


def _pptx_embedded_pictures(data: bytes, *, max_edge: int, max_pages: int) -> list[PageImage]:
    """Fallback when LibreOffice is unavailable: largest picture blob per slide."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(io.BytesIO(data))
    out: list[PageImage] = []
    for index, slide in enumerate(presentation.slides, start=1):
        if index > max_pages:
            break
        blobs = _collect_pictures(slide.shapes, MSO_SHAPE_TYPE.PICTURE)
        if not blobs:
            continue
        blob = max(blobs, key=len)
        try:
            from PIL import Image

            image = Image.open(io.BytesIO(blob))
            out.append(_page_image("slide", index, image, max_edge))
        except Exception:
            logger.warning("pptx_picture_decode_failed", extra={"slide": index})
    return out


def _collect_pictures(shapes: Any, picture_type: Any) -> list[bytes]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    blobs: list[bytes] = []
    for shape in shapes:
        shape_type = getattr(shape, "shape_type", None)
        if shape_type == picture_type:
            try:
                blobs.append(shape.image.blob)
            except Exception:
                continue
        elif shape_type == MSO_SHAPE_TYPE.GROUP:
            try:
                blobs.extend(_collect_pictures(shape.shapes, picture_type))
            except Exception:
                continue
    return blobs
